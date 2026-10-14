#!/usr/bin/env python
"""Convert the HPICDA manuscript Markdown to a submission-style .docx.

Handles headings, bold/italic inline spans, unordered and ordered lists,
horizontal rules and paragraph flow.  Citation markers of the form
"(PMID: 12345678)" are left in place for the author to convert with EndNote.
"""
from __future__ import annotations

import os

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(os.environ.get(
    "HPICDA_ROOT", Path(__file__).resolve().parents[1]))

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_inline(par, text: str, base_bold: bool = False, base_italic: bool = False):
    for piece in INLINE.split(text):
        if not piece:
            continue
        bold, italic, mono = base_bold, base_italic, False
        if piece.startswith("**") and piece.endswith("**"):
            piece, bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            piece, italic = piece[1:-1], True
        elif piece.startswith("`") and piece.endswith("`"):
            piece, mono = piece[1:-1], True
        run = par.add_run(piece)
        run.bold = bold
        run.italic = italic
        if mono:
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def build(md: Path, out: Path, title: str):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5

    for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)

    lines = md.read_text(encoding="utf-8").splitlines()
    i = 0
    in_refs = False
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        s = line.strip()

        if not s:
            i += 1
            continue

        if s in ("---", "***", "___"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pr = p.add_run("_" * 60)
            pr.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1))
            text = m.group(2).strip().rstrip("#").strip()
            if lvl == 1 and not doc.paragraphs:
                # document title
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, f"**{text}**")
                for r in p.runs:
                    r.font.size = Pt(17)
                i += 1
                continue
            if lvl == 1 and not title:
                p = doc.add_heading(level=1)
                add_inline(p, text)
                i += 1
                continue
            h = doc.add_heading(level=min(lvl, 3))
            add_inline(h, text)
            if text.lower().startswith("reference"):
                in_refs = True
            i += 1
            continue

        if re.match(r"^\d+\.\s+", s) and in_refs:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-18)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, s)
            i += 1
            continue

        m = re.match(r"^(\s*)([-*])\s+(.*)$", raw)
        if m:
            indent = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Pt(18 + 18 * indent)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, m.group(3))
            i += 1
            continue

        m = re.match(r"^\s*(\d+)\.\s+(.*)$", raw)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, m.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, s)
        i += 1

    doc.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else \
        ROOT / "25_manuscript_package/HPICDA_manuscript_v1.md"
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    build(src, dst, title="")
